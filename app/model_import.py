"""Phase 2: Semantic Models aus hochgeladenen Dateien (CSV/Excel) erzeugen.

Der Weg ist derselbe, der schon beim KI-Test-Modell funktioniert hat: Die Daten
werden als Power-Query-Inline-Tabelle (#table(...), Import-Modus) direkt in die
Modell-Definition (model.bim / TMSL) eingebettet. Kein externer Connector nötig –
das Modell ist selbst-enthaltend. Nach dem Anlegen lädt ein Refresh die Daten.

Ablauf:
  parse_upload(dateiname, bytes) -> parsed  (Spalten + Typen + Zeilen)
  build_model_bim(name, parsed)  -> model.bim-Dict für ein NEUES Modell
  build_table_object(parsed)     -> eine Tabelle zum Anhängen an ein BESTEHENDES Modell

Grenze: Die Daten liegen in der Definition, daher für PoC/kleine bis mittlere
Dateien gedacht (Zeilen werden auf MAX_ROWS gedeckelt). Große/laufende Quellen
kommen später über echte Connectoren (SharePoint/Addison).
"""
import csv
import datetime as dt
import io
import re
from typing import Any

# Obergrenze eingebetteter Zeilen – hält die Definition handhabbar.
MAX_ROWS = 10000

_DATE_FORMATS = [
    "%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%Y/%m/%d",
    "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S", "%d.%m.%Y %H:%M:%S",
    "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M",
]
_BOOL_TRUE = {"true", "wahr", "ja", "yes"}
_BOOL_FALSE = {"false", "falsch", "nein", "no"}


# ── Werte-Parser (robust gegen deutsche Zahlen-/Datumsformate) ─────────────
def _is_null(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _to_int(v: Any) -> int | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        return int(v) if float(v).is_integer() else None
    if isinstance(v, str):
        s = v.strip()
        if re.fullmatch(r"[+-]?\d+", s):
            return int(s)
        if re.fullmatch(r"[+-]?\d{1,3}(\.\d{3})+", s):   # 1.234.567 (Tausenderpunkte)
            return int(s.replace(".", ""))
    return None


def _to_float(v: Any) -> float | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.strip().replace("€", "").replace("$", "").replace("%", "").replace(" ", "")
        if not s:
            return None
        if "." in s and "," in s:      # 1.234,56  -> de: Punkt=Tausender, Komma=Dezimal
            s = s.replace(".", "").replace(",", ".")
        elif "," in s:                  # 1234,56   -> de: Komma=Dezimal
            s = s.replace(",", ".")
        try:
            return float(s)
        except ValueError:
            return None
    return None


def _to_date(v: Any) -> dt.datetime | None:
    if isinstance(v, dt.datetime):
        return v
    if isinstance(v, dt.date):
        return dt.datetime(v.year, v.month, v.day)
    if isinstance(v, str):
        s = v.strip()
        if s:
            for fmt in _DATE_FORMATS:
                try:
                    return dt.datetime.strptime(s, fmt)
                except ValueError:
                    continue
    return None


def _to_bool(v: Any) -> bool | None:
    if isinstance(v, bool):
        return v
    if isinstance(v, str):
        s = v.strip().lower()
        if s in _BOOL_TRUE:
            return True
        if s in _BOOL_FALSE:
            return False
    return None


def _infer_type(values: list[Any]) -> str:
    """Bestimmt den Spaltentyp: boolean | int64 | double | dateTime | string."""
    non_null = [v for v in values if not _is_null(v)]
    if not non_null:
        return "string"
    if all(_to_bool(v) is not None for v in non_null):
        return "boolean"
    if all(_to_int(v) is not None for v in non_null):
        return "int64"
    if all(_to_float(v) is not None for v in non_null):
        return "double"
    if all(_to_date(v) is not None for v in non_null):
        return "dateTime"
    return "string"


# ── Datei einlesen ─────────────────────────────────────────────────────────
# Python-Kodierung -> Power-Query-Codepage (fürs Csv.Document in Variante B)
_M_ENCODING = {"utf-8-sig": 65001, "utf-8": 65001, "cp1252": 1252, "latin-1": 28591}


def _decode(content: bytes) -> tuple[str, str]:
    """Dekodiert und gibt (text, verwendete_kodierung) zurück."""
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return content.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return content.decode("latin-1", errors="replace"), "latin-1"


def _read_csv(content: bytes) -> tuple[list, list[list], str, str]:
    """Rückgabe: (header, rows, trennzeichen, kodierung)."""
    text, enc = _decode(content)
    sample = text[:4096]
    try:
        delim = csv.Sniffer().sniff(sample, delimiters=";,\t|").delimiter
    except csv.Error:
        delim = ";" if sample.count(";") >= sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    rows = [r for r in reader if any((c or "").strip() for c in r)]
    if not rows:
        raise ValueError("Die CSV-Datei enthält keine Zeilen.")
    return rows[0], rows[1:], delim, enc


def _workbook(content: bytes):
    try:
        from openpyxl import load_workbook
    except ImportError as e:  # noqa: BLE001
        raise RuntimeError("Für Excel-Dateien fehlt die Bibliothek 'openpyxl'.") from e
    return load_workbook(io.BytesIO(content), read_only=True, data_only=True)


def list_sheets(content: bytes) -> list[str]:
    """Blattnamen einer Excel-Datei – Grundlage für die Blattauswahl."""
    wb = _workbook(content)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _read_xlsx(content: bytes, sheet: str | None = None) -> tuple[list, list[list]]:
    """Liest ein Blatt. Ohne Angabe das aktive – das ist bei Mehrblatt-Dateien
    oft das falsche (z. B. ein leeres Trennblatt), deshalb sollte 'sheet' gesetzt
    werden, sobald die Datei mehrere Blätter hat."""
    wb = _workbook(content)
    if sheet:
        if sheet not in wb.sheetnames:
            wb.close()
            raise ValueError(f"Blatt '{sheet}' gibt es nicht. Vorhanden: {', '.join(wb.sheetnames)}")
        ws = wb[sheet]
    else:
        ws = wb.active
    header: list | None = None
    data: list[list] = []
    for row in ws.iter_rows(values_only=True):
        if row is None or all(c is None for c in row):
            continue
        if header is None:
            header = list(row)
        else:
            data.append(list(row))
    wb.close()
    if header is None:
        raise ValueError("Das Excel-Blatt ist leer.")
    return header, data


def _clean_headers(header: list) -> list[str]:
    out: list[str] = []
    used: set[str] = set()
    for i, h in enumerate(header):
        name = (str(h).strip() if h is not None else "") or f"Spalte_{i + 1}"
        cand, k = name, 1
        while cand.lower() in used:
            k += 1
            cand = f"{name}_{k}"
        used.add(cand.lower())
        out.append(cand)
    return out


def sanitize_name(name: str) -> str:
    """Bereinigt Modell-/Tabellennamen (Umlaute erlaubt, Sonderzeichen raus)."""
    name = re.sub(r"[^\w äöüÄÖÜß.\-]", " ", name or "", flags=re.UNICODE)
    name = re.sub(r"\s+", " ", name).strip()
    return name[:80] or "Tabelle"


def parse_upload(filename: str, content: bytes,
                 sheet: str | None = None) -> dict[str, Any]:
    """Liest CSV/Excel, erkennt Spalten & Typen, liefert die Zeilen.

    'sheet' wählt bei Excel das Blatt (nötig bei Mehrblatt-Dateien).
    Zusätzlich für Variante B: 'delimiter' und 'encoding' (Codepage) der Datei –
    damit Power Query dieselbe Datei später genauso liest wie wir hier.
    """
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    is_excel = ext in ("xlsx", "xlsm")
    if is_excel:
        header, rows = _read_xlsx(content, sheet)
        delim, enc = None, None
    else:
        header, rows, delim, enc = _read_csv(content)

    header = _clean_headers(header)
    truncated = len(rows) > MAX_ROWS
    rows = rows[:MAX_ROWS]

    columns: list[dict[str, Any]] = []
    for i, name in enumerate(header):
        vals = [r[i] if i < len(r) else None for r in rows]
        dtype = _infer_type(vals)
        has_time = dtype == "dateTime" and any(
            (d := _to_date(v)) and (d.hour or d.minute or d.second)
            for v in vals if not _is_null(v)
        )
        columns.append({"name": name, "dtype": dtype, "has_time": bool(has_time)})

    base = filename.rsplit(".", 1)[0] if "." in filename else filename
    return {
        "table_name": sanitize_name(base),
        "columns": columns,
        "rows": rows,
        "row_count": len(rows),
        "truncated": truncated,
        "is_excel": is_excel,
        "delimiter": delim,
        "encoding": _M_ENCODING.get(enc or "", 65001),
    }


# ── model.bim / TMSL aufbauen ──────────────────────────────────────────────
_M_TYPE = {"int64": "Int64.Type", "double": "number", "boolean": "logical",
           "string": "text"}
_TMSL_TYPE = {"int64": "int64", "double": "double", "boolean": "boolean",
              "dateTime": "dateTime", "string": "string"}


def _m_type(dtype: str, has_time: bool) -> str:
    if dtype == "dateTime":
        return "datetime" if has_time else "date"
    return _M_TYPE[dtype]


def _m_value(v: Any, dtype: str, has_time: bool) -> str:
    if _is_null(v):
        return "null"
    if dtype == "int64":
        i = _to_int(v)
        return "null" if i is None else str(i)
    if dtype == "double":
        f = _to_float(v)
        return "null" if f is None else repr(f)
    if dtype == "boolean":
        b = _to_bool(v)
        return "null" if b is None else ("true" if b else "false")
    if dtype == "dateTime":
        d = _to_date(v)
        if d is None:
            return "null"
        if has_time:
            return f"#datetime({d.year},{d.month},{d.day},{d.hour},{d.minute},{d.second})"
        return f"#date({d.year},{d.month},{d.day})"
    # string
    s = str(v).replace('"', '""').replace("\r", " ").replace("\n", " ")
    return f'"{s}"'


def _tmsl_column(name: str, dtype: str) -> dict[str, Any]:
    col: dict[str, Any] = {
        "name": name,
        "dataType": _TMSL_TYPE[dtype],
        "sourceColumn": name,
        "summarizeBy": "sum" if dtype in ("int64", "double") else "none",
    }
    if dtype == "dateTime":
        col["formatString"] = "General Date"
    return col


def _build_m(columns: list[dict[str, Any]], rows: list[list]) -> list[str]:
    type_decl = ", ".join(
        f'#"{c["name"]}" = {_m_type(c["dtype"], c["has_time"])}' for c in columns)
    row_lines = []
    for r in rows:
        vals = [_m_value(r[i] if i < len(r) else None, c["dtype"], c["has_time"])
                for i, c in enumerate(columns)]
        row_lines.append("            {" + ", ".join(vals) + "}")
    return [
        "let",
        "    Source = #table(",
        f"        type table [{type_decl}],",
        "        {",
        ",\n".join(row_lines),
        "        }",
        "    )",
        "in",
        "    Source",
    ]


def build_table_object(parsed: dict[str, Any]) -> dict[str, Any]:
    """Baut ein TMSL-Tabellenobjekt (Spalten + Inline-Import-Partition)."""
    name = parsed["table_name"]
    columns = parsed["columns"]
    return {
        "name": name,
        "columns": [_tmsl_column(c["name"], c["dtype"]) for c in columns],
        "partitions": [{
            "name": name,
            "mode": "import",
            "source": {"type": "m", "expression": _build_m(columns, parsed["rows"])},
        }],
    }


def build_relationship(from_table: str, from_column: str,
                       to_table: str, to_column: str,
                       is_active: bool = True) -> dict[str, Any]:
    """TMSL-Beziehung (viele-zu-eins; das ist die TMSL-Vorgabe, wenn nichts anderes
    angegeben ist). 'to_column' muss eindeutig sein, sonst lädt das Modell nicht.

    is_active=False legt die Beziehung inaktiv an (isActive=false) – nötig, wenn sie
    sonst einen zweiten Pfad zwischen bereits verbundenen Tabellen erzeugen würde
    (Power BI lehnt mehrdeutige Pfade beim Anlegen ab)."""
    rel: dict[str, Any] = {
        "name": f"{from_table}_{from_column}__{to_table}_{to_column}"[:100],
        "fromTable": from_table,
        "fromColumn": from_column,
        "toTable": to_table,
        "toColumn": to_column,
    }
    if not is_active:
        rel["isActive"] = False
    return rel


def build_fk_relationships(fks: list[dict[str, str]],
                           table_names: list[str]) -> list[dict[str, Any]]:
    """Baut TMSL-Beziehungen 1:1 aus den in der DB deklarierten Foreign Keys.

    Kein KI-Raten, keine Daten- oder Schema-Weitergabe an ein LLM: Die Beziehungen
    kommen direkt aus mssql.foreign_keys(). Berücksichtigt werden nur FKs, deren
    beide Tabellen im Modell liegen ('table_names').

    Union-Find-Wächter: Würde ein FK eine ZWEITE Verbindung zwischen bereits
    (indirekt) verbundenen Tabellen bilden, wird er inaktiv angelegt statt
    verworfen – so bleibt das Schema vollständig abgebildet, ohne dass Power BI
    mehrdeutige Pfade (ambiguous paths) bekommt. Beim aktuellen Addison-Schema
    entstehen 0 inaktive; der Wächter ist Absicherung gegen künftige Änderungen.
    """
    names = set(table_names)
    parent: dict[str, str] = {}

    def root(x: str) -> str:
        parent.setdefault(x, x)
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    # Dimensions-FKs zuerst, Mandanten-FKs (-> StammMandant) zuletzt einordnen:
    # entstünde je ein Zyklus, würde eher eine redundante Mandanten-Beziehung
    # inaktiv als eine fachliche.
    ordered = sorted(fks, key=lambda f: (f.get("ref_table") == "StammMandant",
                                         f.get("parent_table", ""), f.get("parent_col", "")))
    out: list[dict[str, Any]] = []
    seen: set[tuple] = set()
    for f in ordered:
        ft, fc = f.get("parent_table"), f.get("parent_col")
        tt, tc = f.get("ref_table"), f.get("ref_col")
        if ft not in names or tt not in names or ft == tt:
            continue
        key = (ft, fc, tt, tc)
        if key in seen:
            continue
        seen.add(key)
        active = root(ft) != root(tt)
        if active:
            parent[root(ft)] = root(tt)
        out.append(build_relationship(ft, fc, tt, tc, is_active=active))
    return out


def build_model_bim_tables(model_name: str, table_objs: list[dict[str, Any]],
                           relationships: list[dict[str, Any]] | None = None) -> dict[str, Any]:
    """Baut ein model.bim für ein neues Modell mit einer oder mehreren Tabellen."""
    model: dict[str, Any] = {
        "culture": "de-DE",
        "defaultPowerBIDataSourceVersion": "powerBI_V3",
        "tables": table_objs,
    }
    if relationships:
        model["relationships"] = relationships
    return {"name": model_name, "compatibilityLevel": 1600, "model": model}


def build_model_bim(model_name: str, parsed: dict[str, Any]) -> dict[str, Any]:
    """Baut ein komplettes model.bim für ein neues Semantic Model (eine Tabelle)."""
    return build_model_bim_tables(model_name, [build_table_object(parsed)])


def merge_parsed(parsed_list: list[dict[str, Any]],
                 table_name: str | None = None) -> dict[str, Any]:
    """Führt mehrere gleich strukturierte Dateien zu EINER Tabelle zusammen.

    Voraussetzung: identische Spaltennamen in identischer Reihenfolge (dieselbe
    Struktur, wie sie in einem Ordner liegen soll). Bei Abweichung: klarer Fehler.
    Typen werden über die kombinierten Zeilen neu bestimmt (robuster als je Datei).
    """
    if not parsed_list:
        raise ValueError("Keine Dateien zum Zusammenführen.")
    base_cols = [c["name"] for c in parsed_list[0]["columns"]]
    rows: list[list] = []
    truncated = False
    for p in parsed_list:
        cols = [c["name"] for c in p["columns"]]
        if cols != base_cols:
            raise ValueError(
                f"Spalten passen nicht zusammen ({cols} ≠ {base_cols}). "
                "Nur Dateien gleicher Struktur können an eine Tabelle angehängt werden.")
        rows.extend(p["rows"])
        truncated = truncated or p["truncated"]
    truncated = truncated or len(rows) > MAX_ROWS
    rows = rows[:MAX_ROWS]

    columns: list[dict[str, Any]] = []
    for i, name in enumerate(base_cols):
        vals = [r[i] if i < len(r) else None for r in rows]
        dtype = _infer_type(vals)
        has_time = dtype == "dateTime" and any(
            (d := _to_date(v)) and (d.hour or d.minute or d.second)
            for v in vals if not _is_null(v))
        columns.append({"name": name, "dtype": dtype, "has_time": bool(has_time)})

    first = parsed_list[0]
    return {
        "table_name": sanitize_name(table_name or first["table_name"]),
        "columns": columns,
        "rows": rows,
        "row_count": len(rows),
        "truncated": truncated,
        # Datei-Eigenschaften der ersten Datei durchreichen (Variante B braucht sie)
        "is_excel": first.get("is_excel", False),
        "delimiter": first.get("delimiter"),
        "encoding": first.get("encoding", 65001),
    }


# ── Variante B: echte SharePoint-Quelle statt eingebetteter Daten ──────────
# Die Daten stehen NICHT im Modell – Power BI holt sie beim Refresh selbst.
# Damit gibt es keine Zeilengrenze, aber das Dataset braucht einmalig
# SharePoint-Anmeldedaten (siehe /sharepoint/create-model).

def _esc(s: str) -> str:
    """String für ein M-Literal in doppelten Anführungszeichen escapen."""
    return str(s).replace('"', '""')


def _m_delim(d: str) -> str:
    """Trennzeichen als M-Literal (Tab braucht die #(tab)-Schreibweise)."""
    return {"\t": "#(tab)"}.get(d, _esc(d))


def _m_type_b(dtype: str, has_time: bool) -> str:
    """M-Typ für Table.TransformColumnTypes."""
    if dtype == "int64":
        return "Int64.Type"
    if dtype == "double":
        return "type number"
    if dtype == "boolean":
        return "type logical"
    if dtype == "dateTime":
        return "type datetime" if has_time else "type date"
    return "type text"


def build_sharepoint_m(site_url: str, folder: str, files: list[str],
                       columns: list[dict[str, Any]], delimiter: str | None,
                       encoding: int, is_excel: bool,
                       recursive: bool = False, extension: str = ".csv",
                       sheet: str | None = None) -> list[str]:
    """Baut die M-Abfrage, die Power BI zu den Dateien in SharePoint schickt.

    recursive=True: alle Dateien UNTERHALB des Ordners (Filter über den Pfad).
    Damit fließen später hinzugefügte Dateien – etwa ein neuer Monatsordner –
    beim nächsten Refresh automatisch mit, ohne dass das Modell angefasst wird.
    recursive=False: nur die namentlich genannten Dateien.

    Wichtig: Header werden PRO DATEI hochgestuft und erst danach kombiniert –
    sonst landen die Kopfzeilen der Folgedateien als Datenzeilen in der Tabelle.
    Die Typumwandlung läuft mit Kultur "de-DE" (deutsche Zahlen-/Datumsformate).
    """
    folder = (folder or "").strip("/")
    if recursive:
        cond = f'Text.Lower([Extension]) = "{_esc(extension.lower())}"'
        if folder:  # auf den Teilbaum eingrenzen; "/" im Muster trennt Namensteile
            cond += f' and Text.Contains([Folder Path], "/{_esc(folder)}/")'
    else:
        namelist = ", ".join(f'"{_esc(f)}"' for f in files)
        cond = f"List.Contains({{{namelist}}}, [Name])"
        if folder:  # gleiche Dateinamen könnten anderswo liegen
            seg = folder.rsplit("/", 1)[-1]
            cond += f' and Text.EndsWith([Folder Path], "/{_esc(seg)}/")'

    if is_excel:
        # Blatt per NAME ansprechen, nicht per Position: {0} waere das erste Blatt
        # und damit falsch, sobald jemand Blätter umsortiert oder eines einfügt.
        nav = (f'{{[Item="{_esc(sheet)}",Kind="Sheet"]}}' if sheet else "{0}")
        read = f"Excel.Workbook([Content], true){nav}[Data]"
    else:
        read = (f'Table.PromoteHeaders(Csv.Document([Content], '
                f'[Delimiter="{_m_delim(delimiter or ";")}", Encoding={encoding}, '
                f'QuoteStyle=QuoteStyle.Csv]), [PromoteAllScalars=true])')

    types = ", ".join(f'{{"{_esc(c["name"])}", {_m_type_b(c["dtype"], c["has_time"])}}}'
                      for c in columns)
    # Spalten POSITIONSBASIERT auf unsere Namen umbenennen, BEVOR Typen gesetzt
    # werden. Wichtig bei Quelldateien mit leerer Kopfzeile: Power BI nennt eine
    # solche Spalte z. B. "Column3", unser Parser aber "Spalte_3" – ohne dieses
    # Mapping würde TransformColumnTypes die Spalte nicht finden (Refresh scheitert).
    names = ", ".join(f'"{_esc(c["name"])}"' for c in columns)
    return [
        "let",
        f'    Quelle = SharePoint.Files("{_esc(site_url)}", [ApiVersion = 15]),',
        f"    Auswahl = Table.SelectRows(Quelle, each {cond}),",
        f'    MitDaten = Table.AddColumn(Auswahl, "Daten", each {read}),',
        "    Kombiniert = Table.Combine(MitDaten[Daten]),",
        f'    Benannt = Table.RenameColumns(Kombiniert, '
        f'List.Zip({{Table.ColumnNames(Kombiniert), {{{names}}}}}), MissingField.Ignore),',
        f'    Typen = Table.TransformColumnTypes(Benannt, {{{types}}}, "de-DE")',
        "in",
        "    Typen",
    ]


def build_sharepoint_table_object(parsed: dict[str, Any], table_name: str,
                                  site_url: str, folder: str, files: list[str],
                                  recursive: bool = False, extension: str = ".csv",
                                  sheet: str | None = None) -> dict[str, Any]:
    """TMSL-Tabelle, deren Partition live auf SharePoint zeigt (Variante B)."""
    name = sanitize_name(table_name or parsed["table_name"])
    columns = parsed["columns"]
    return {
        "name": name,
        "columns": [_tmsl_column(c["name"], c["dtype"]) for c in columns],
        "partitions": [{
            "name": name,
            "mode": "import",
            "source": {"type": "m",
                       "expression": build_sharepoint_m(
                           site_url, folder, files, columns,
                           parsed.get("delimiter"), parsed.get("encoding", 65001),
                           parsed.get("is_excel", False), recursive, extension,
                           sheet)},
        }],
    }


# ── Variante C: MS SQL (Addison) als Live-Quelle ───────────────
# Wie Variante B holt Power BI die Daten beim Refresh selbst – nur über
# Sql.Database() statt SharePoint.Files(). Keine Zeilengrenze, keine
# eingebetteten Daten. Die Typen liefert SQL Server nativ mit, deshalb ist
# die M-Abfrage schlanker als bei CSV (kein locale-Parsing nötig).

# SQL-Server-Datentyp -> TMSL-Datentyp. Alles Unbekannte wird 'string'.
_SQL_TO_TMSL = {
    "bigint": "int64", "int": "int64", "smallint": "int64", "tinyint": "int64",
    "bit": "boolean",
    "float": "double", "real": "double", "decimal": "double", "numeric": "double",
    "money": "double", "smallmoney": "double",
    "datetime": "dateTime", "datetime2": "dateTime", "smalldatetime": "dateTime",
    "date": "dateTime", "datetimeoffset": "dateTime",
}


def _sql_tmsl_type(sql_type: str) -> str:
    return _SQL_TO_TMSL.get((sql_type or "").lower(), "string")


def _sql_tmsl_column(name: str, sql_type: str) -> dict[str, Any]:
    tmsl = _sql_tmsl_type(sql_type)
    col: dict[str, Any] = {
        "name": name,
        "dataType": tmsl,
        "sourceColumn": name,
        "summarizeBy": "sum" if tmsl in ("int64", "double") else "none",
    }
    if tmsl == "dateTime":
        col["formatString"] = "General Date"
    return col


# SQL-Datumstypen -> (M-Untergrenze, M-Ergebnistyp) für die Datumsbereinigung.
_SQL_DATE_KIND = {
    "date": ("#date(1900, 3, 1)", "type nullable date"),
    "datetime": ("#datetime(1900, 3, 1, 0, 0, 0)", "type nullable datetime"),
    "datetime2": ("#datetime(1900, 3, 1, 0, 0, 0)", "type nullable datetime"),
    "smalldatetime": ("#datetime(1900, 3, 1, 0, 0, 0)", "type nullable datetime"),
}


def _sql_date_transforms(columns: list[dict[str, str]]) -> str:
    """M-Transformationsliste, die in JEDER Datumsspalte Werte vor dem 01.03.1900
    auf null setzt.

    Hintergrund: Power BI/VertiPaq unterstützt Datumswerte erst ab dem 01.03.1900.
    Ältere Sentinel-/Leerwerte (in Buchhaltungsdaten üblich, z. B. 1900-01-01 oder
    0001-01-01) lösen beim Refresh sonst 'invalid date'-Fehler aus. Ausgewertet
    werden nur die Spalten-TYPEN (Metadaten) – KEINE Datenwerte; die eigentliche
    Bereinigung führt Power BI beim Laden aus.
    """
    parts: list[str] = []
    for c in columns:
        kind = _SQL_DATE_KIND.get((c.get("type") or "").lower())
        if not kind:
            continue
        bound, rtype = kind
        parts.append(
            f'{{"{_esc(c["name"])}", each if _ = null then null '
            f'else if _ < {bound} then null else _, {rtype}}}')
    return ", ".join(parts)


def build_sql_m(server: str, port: int, database: str, schema: str, table: str,
                mandant_col: str | None = None, mandant_val: str | None = None,
                columns: list[dict[str, str]] | None = None) -> list[str]:
    """M-Abfrage, die Power BI an den SQL-Server schickt (eine Tabelle).

    mandant_col/mandant_val filtern optional auf einen Mandanten – nötig, solange
    wir Single-Mandant-Modelle bauen (dann ist z. B. KONTO eindeutig und
    Beziehungen funktionieren ohne kombinierten Schlüssel).

    'columns' (mssql.get_columns()): werden mitgegeben, um Datumsspalten gegen
    ungültige Datümer (< 01.03.1900) abzusichern (siehe _sql_date_transforms).
    """
    srv = server if int(port) == 1433 else f"{server},{port}"
    lines = [
        "let",
        f'    Quelle = Sql.Database("{_esc(srv)}", "{_esc(database)}"),',
        f'    Navigation = Quelle{{[Schema="{_esc(schema)}", Item="{_esc(table)}"]}}[Data]',
    ]
    final = "Navigation"
    if mandant_col and mandant_val is not None:
        lines[-1] += ","
        lines.append(
            f'    Gefiltert = Table.SelectRows(Navigation, each [{_esc(mandant_col)}] = '
            f'"{_esc(mandant_val)}")')
        final = "Gefiltert"
    date_tr = _sql_date_transforms(columns or [])
    if date_tr:
        lines[-1] += ","
        lines.append(f"    Bereinigt = Table.TransformColumns({final}, {{{date_tr}}})")
        final = "Bereinigt"
    lines += ["in", f"    {final}"]
    return lines


def build_sql_table_object(server: str, port: int, database: str, schema: str,
                           table: str, columns: list[dict[str, str]],
                           table_name: str | None = None,
                           mandant_col: str | None = None,
                           mandant_val: str | None = None) -> dict[str, Any]:
    """TMSL-Tabelle, deren Partition live auf eine SQL-Server-Tabelle zeigt.

    'columns' ist die Ausgabe von mssql.get_columns() ([{name, type}]) – der
    SQL-Datentyp wird auf den passenden TMSL-Typ abgebildet.
    """
    name = sanitize_name(table_name or table)
    return {
        "name": name,
        "columns": [_sql_tmsl_column(c["name"], c["type"]) for c in columns],
        "partitions": [{
            "name": name,
            "mode": "import",
            "source": {"type": "m",
                       "expression": build_sql_m(server, port, database, schema,
                                                 table, mandant_col, mandant_val,
                                                 columns=columns)},
        }],
    }


def display_value(v: Any) -> str:
    """Wert für die Vorschau-Anzeige JSON-sicher in Text wandeln."""
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat(sep=" ")[:19]
    return str(v)
